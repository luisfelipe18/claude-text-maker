# --- filepath: video_narrative_processor/core/processors/document_generator.py
import json, re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from core.processors.base import WordGenerator
from core.models.artifacts import WordDocument
from utils.file_utils import sanitize_filename

class DefaultWordGenerator(WordGenerator):
    def build(self, rewritten_json: Path, out_dir: Path, serial: int) -> WordDocument:
        data = json.loads(rewritten_json.read_text(encoding="utf-8"))
        body = (data.get("cuerpo") or "").strip()
        title = (data.get("titulo") or "").strip()
        safe = sanitize_filename(title) or "sin-titulo"
        out_docx = out_dir / f"video_{serial:03d}_{safe}.docx"

        doc = Document()
        style = doc.styles["Normal"]
        font = style.font; font.name = "Calibri"; font.size = Pt(12)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        blocks = re.split(r"\r?\n\s*\r?\n", body)
        for i, b in enumerate(blocks):
            p = doc.add_paragraph(b.strip())
            p.paragraph_format.line_spacing = 1.15
            if i < len(blocks) - 1:
                doc.add_paragraph("")
        out_dir.mkdir(parents=True, exist_ok=True)
        doc.save(out_docx)
        return WordDocument(docx_path=out_docx)